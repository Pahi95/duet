"""Build V4 Word files only from publication/v4 (no original expression data needed).

python scripts/build_v4_documents.py --output ../manuscript
Optional --redline-against PATH creates a third copy marking new words in red.
Requires python-docx and pandas. The original documents are never overwritten.
"""
import argparse, difflib, json, re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from mendeley_v4 import restore as restore_mendeley


def fresh(landscape=False):
    d=Document()
    s=d.sections[0]
    s.page_width=Inches(11.69 if landscape else 8.27)
    s.page_height=Inches(8.27 if landscape else 11.69)
    s.top_margin=s.bottom_margin=Inches(.75)
    s.left_margin=s.right_margin=Inches(.65)
    for name in ['Normal','Heading 1','Heading 2','Title']:
        st=d.styles[name]
        st.font.name='Times New Roman'
        st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11 if name=='Normal' else 13)
        st.paragraph_format.line_spacing=2 if not landscape else 1.15
        st.paragraph_format.space_after=Pt(5)
        for border in st.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        fonts=st.element.find(qn('w:rPr')).find(qn('w:rFonts'))
        if fonts is not None:
            for attr in list(fonts.attrib):
                if attr.endswith('Theme'):del fonts.attrib[attr]
    ln=OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'),'1'); ln.set(qn('w:restart'),'continuous')
    if not landscape: s._sectPr.append(ln)
    p=s.footer.paragraphs[0]; p.alignment=1
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    d.core_properties.author=''
    d.core_properties.last_modified_by=''
    d.core_properties.title='DUET manuscript V4' if not landscape else 'DUET supplementary tables V4'
    return d


def inline(p,text):
    for bit in re.split(r'(\*\*.*?\*\*|==.*?==|`.*?`)',text):
        r=p.add_run(bit[2:-2] if bit.startswith(('**','==')) else bit.strip('`'))
        if bit.startswith('**'):r.bold=True
        if bit.startswith('=='):r.font.highlight_color=WD_COLOR_INDEX.YELLOW
        if bit.startswith('`'):r.font.name='Consolas'


def table(d,caption,df,landscape=False):
    # Keep every column in one table. Rows may continue with a repeated header.
    p=d.add_paragraph();inline(p,caption)
    p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.1
    s=d.sections[-1];available=(s.page_width-s.left_margin-s.right_margin)/914400
    rendered=[['' if pd.isna(v) else (f'{v:.4g}' if isinstance(v,float) else str(v)) for v in row] for row in df.itertuples(index=False,name=None)]
    grouped='wrapper_vs_pydeseq2_n_genes' in df.columns
    labels=[str(c).replace('_',' ') for c in df.columns]
    if grouped:
        metrics=['Genes','log2FC\nr','log2FC\nmax |Δ|','−log10 p\nρ','−log10 p\nmax |Δ|','DE\nA','DE\nB','DE\nJaccard']
        labels=['Dataset','Cell type','Design','n','n\nref.','n\ntest','Identical\nsums']+metrics+metrics+['Identical\nwithin\n10⁻⁸']
        inline(p,' r: Pearson correlation; ρ: Spearman correlation; DE: number of significant genes; Δ: difference between implementations.')
    weights=[]
    for j,col in enumerate(df.columns):
        values=[labels[j]]+[r[j] for r in rendered]
        longest=max(len(word) for value in values for word in value.split())
        mean=sum(min(len(v),60) for v in values)/len(values)
        weights.append(max(6,min(23,longest*.7+mean*.2)))
    widths=[available*w/sum(weights) for w in weights]
    if grouped:
        weights=[.65,1.15,.85,.4,.4,.4,.6]+[.52,.55,.85,.55,.65,.55,.55,.6]*2+[.8]
        widths=[available*w/sum(weights) for w in weights]
    t=d.add_table(rows=2 if grouped else 1,cols=len(df.columns));t.style='Table Grid';t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    header_count=2 if grouped else 1
    for row in t.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell,label in zip(t.rows[header_count-1].cells,labels):cell.text=label
    for values in rendered:
        row=t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,value in zip(row.cells,values):cell.text=value
    for ri,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if grouped:
                margins=OxmlElement('w:tcMar')
                for side in ['left','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'50');edge.set(qn('w:type'),'dxa');margins.append(edge)
                cell._tc.get_or_add_tcPr().append(margins)
            if ri<header_count:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1
                padding=2 if len(df)>40 else 3
                p.paragraph_format.space_after=Pt(padding);p.paragraph_format.space_before=Pt(padding)
                for r in p.runs:r.font.size=Pt(9);r.bold=ri<header_count
    if grouped:
        for start,end,title in [(0,6,'Dataset and design'),(7,14,'A: DUET wrapper; B: direct PyDESeq2'),(15,22,'A: PyDESeq2; B: R DESeq2'),(23,23,'Wrapper /\nPyDESeq2')]:
            cell=t.rows[0].cells[start]
            if end>start:cell=cell.merge(t.rows[0].cells[end])
            cell.text=title
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9);r.bold=True
    d.add_paragraph().paragraph_format.space_after=Pt(3)


def section(d,landscape=False,a3=False):
    s=d.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width=Inches(16.54 if a3 else 11.69 if landscape else 8.27)
    s.page_height=Inches(11.69 if a3 else 8.27 if landscape else 11.69)
    s.left_margin=s.right_margin=Inches(.6)
    s.top_margin=s.bottom_margin=Inches(.55 if landscape else .75)
    for e in list(s._sectPr):
        if e.tag==qn('w:lnNumType'):s._sectPr.remove(e)
    if not landscape:
        ln=OxmlElement('w:lnNumType');ln.set(qn('w:countBy'),'1');ln.set(qn('w:restart'),'continuous');s._sectPr.append(ln)
    return s


def vancouver(s):
    m=re.fullmatch(r'\[(\d+)\](.*?)“(.*?),” (.*?), vol\. (.*?), (?:no\. (.*?), )?pp?\. (.*?), (\d{4}), doi: (.*?)\.',s)
    if not m:return s
    n,authors,title,journal,vol,issue,pages,year,doi=m.groups()
    aa=[]
    for a in authors.strip(' ,').replace(', and ', ', ').replace(' and ', ', ').split(', '):
        a=a.removeprefix('and ')
        suffix=' et al.' if ' et al.' in a else ''
        a=a.replace(' et al.','')
        am=re.match(r'^((?:(?:[A-Z][.\- ]+)|(?:[A-Z] ))+)(.+)$',a)
        aa.append((am[2]+' '+re.sub(r'[. ]','',am[1]) if am else a)+suffix)
    return f'[{n}] '+', '.join(aa)+f'. {title}. {journal} {year};{vol}'+(f'({issue})' if issue else '')+f':{pages}. doi:{doi}.'


def build(source,output,redline=None):
    output.mkdir(parents=True,exist_ok=True)
    meta=json.loads((source/'tables.json').read_text(encoding='utf-8'))
    tabs={t['id']:t for t in meta}
    doc=fresh()
    pending_tables=[];heatmap=False
    for line in (source/'manuscript.md').read_text(encoding='utf-8').splitlines():
        if not line.strip():continue
        if line.startswith('@@TABLE'):
            pending_tables.append(tabs[line.split()[1]]);continue
        if line.startswith('@@FIGURE'):
            heatmap=line.split()[1]=='fig3_effects.png'
            if heatmap:section(doc,True)
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1
            p.add_run().add_picture(str(source/'figures'/line.split()[1]),width=Inches(10.45 if heatmap else 6.9))
            # Figures may be followed by long captions: let captions flow instead of forcing blank pages.
            continue
        if line.startswith('@@REFERENCES'):
            # A real Mendeley bibliography control, retaining the author's style.
            sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
            tag=OxmlElement('w:tag');tag.set(qn('w:val'),'MENDELEY_BIBLIOGRAPHY');pr.append(tag)
            ident=OxmlElement('w:id');ident.set(qn('w:val'),'1771052620');pr.append(ident);sdt.append(pr)
            content=OxmlElement('w:sdtContent');sdt.append(content)
            for ref in (source/'references.txt').read_text(encoding='utf-8').splitlines():
                p=doc.add_paragraph(re.sub(r'^(\[\d+\])',r'\1 ',ref));content.append(p._p)
            doc._body._body.insert(-1,sdt)
            continue
        if line=='## References' and heatmap:section(doc);heatmap=False
        style='Normal'
        for prefix,st in [('### ','Heading 2'),('## ','Heading 1'),('# ','Title')]:
            if line.startswith(prefix):line=line[len(prefix):];style=st;break
        p=doc.add_paragraph(style=style);inline(p,line.removeprefix('- '))
        if heatmap:p.paragraph_format.line_spacing=1.1
    for t in pending_tables:
        section(doc,True)
        table(doc,t['caption'],pd.read_csv(source/t['file']),True)
    main=output/'DUET_manuscript_revised_MO_v4.docx';doc.save(main)
    sup=fresh(True)
    first=sup.sections[0]
    first.page_width=Inches(16.54);first.page_height=Inches(11.69)
    first.orientation=WD_ORIENT.LANDSCAPE
    first.left_margin=first.right_margin=Inches(.6)
    first.top_margin=first.bottom_margin=Inches(.55)
    sup.add_paragraph('Additional file 2. Supplementary tables — V4',style='Title')
    sup.add_paragraph('Machine-readable CSVs accompany every table. Memory fields ending in _mb in archived source files are measured in MiB (bytes / 2²⁰).')
    first_table=True
    for t in meta:
        if t['supplement']:
            df=pd.read_csv(source/t['file'])
            if not first_table:section(sup,True,len(df.columns)>12 or len(df)>26)
            first_table=False
            table(sup,t['caption'],df,True)
    sup.save(output/'DUET_additional_file_2_tables_v4.docx')
    if redline:
        old=Document(redline)
        oldpars=[p.text for p in old.paragraphs if p.text.strip()]
        for p in doc.paragraphs:
            if not p.text.strip() or not p.runs:continue
            candidates=sorted(oldpars,key=lambda x:difflib.SequenceMatcher(None,x,p.text).quick_ratio(),reverse=True)[:3]
            best=max(candidates,key=lambda x:difflib.SequenceMatcher(None,x.split(),p.text.split()).ratio())
            new=p.text
            if new==best:continue
            # Word-token diff preserves equal text and marks only additions/replacements red.
            a=re.findall(r'\S+\s*',best);b=re.findall(r'\S+\s*',new)
            p.clear()
            for tag,_,_,j1,j2 in difflib.SequenceMatcher(None,a,b,autojunk=False).get_opcodes():
                if tag=='delete':continue
                r=p.add_run(''.join(b[j1:j2]))
                if tag!='equal':r.font.color.rgb=RGBColor(192,0,0)
        doc.save(output/'DUET_manuscript_revised_MO_v4_red.docx')
        restore_mendeley(output/'DUET_manuscript_revised_MO_v4_red.docx',source)
    restore_mendeley(main,source)
    print('Wrote V4 manuscript, supplementary tables'+(' and additions-marked copy.' if redline else '.'))


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source',type=Path,default=Path(__file__).resolve().parents[1]/'publication/v4')
    ap.add_argument('--output',type=Path,required=True)
    ap.add_argument('--redline-against',type=Path)
    a=ap.parse_args();build(a.source,a.output,a.redline_against)
